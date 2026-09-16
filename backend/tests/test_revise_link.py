"""统一链路（风险预警 / 条款比对 → 修改合同）的后端行为测试。

覆盖本轮新增/修复的能力：
1. 条款比对来源（clause_key 非 AuditRecord.id）→ `/revise` 用 `_locate_clause` 兜底生成 DOCX 原文锚点
2. 比对会话多轮 → 沿用上一轮锚点，DOCX 采用最新一轮结果
3. `/revised-docx` 存在「无锚点」修订时明确 400，不再静默交付漏改的修订版
4. 风险来源锚点仍取自审核阶段 AuditRecord（未被兜底逻辑改变）

运行（backend 目录下）：
    python -m unittest tests.test_revise_link -v
"""
import os
import sys
import json
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

_BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))
os.environ.setdefault("SECRET_KEY", "test-secret-key-not-weak-123456")

from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from fastapi import HTTPException, BackgroundTasks  # noqa: E402

from database import Base  # noqa: E402
from models.contract import Contract  # noqa: E402
from models.audit_record import AuditRecord  # noqa: E402
from models.clause_revision import ClauseRevision  # noqa: E402
import api.contracts as contracts  # noqa: E402
import ai.reviser as reviser  # noqa: E402

DOC_PARAS = [
    "第一条 合同标的与范围",
    "第二条 验收标准与验收方式：本合同项下的验收采用人工审核方式，验收合格后30日内组织验收。",
    "第三条 违约责任与违约金上限为合同总价的30%。",
]
PARSED = "\n".join(DOC_PARAS)
MATCHED = "验收标准与验收方式：本合同项下的验收采用人工审核方式"
CMP_KEY = "__cmp__验收标准"
RISK_CLAUSE = "违约责任与违约金上限为合同总价的30%"
RISK_ANCHOR = "第三条 违约责任与违约金上限为合同总价的30%。"


class TestUnifiedReviseLink(unittest.TestCase):
    """条款比对来源的锚点兜底 + DOCX 漏改显式失败。"""

    def _setup(self):
        tmp = tempfile.TemporaryDirectory()
        eng = create_engine(f"sqlite:///{Path(tmp.name) / 't.db'}", connect_args={"timeout": 30})
        Base.metadata.create_all(eng)
        self.uploads = Path(tmp.name) / "uploads"
        self.uploads.mkdir(exist_ok=True)
        return tmp, eng, sessionmaker(bind=eng)

    def _make_docx(self, path):
        doc = Document()
        for p in DOC_PARAS:
            doc.add_paragraph(p)
        doc.save(str(path))
        return str(path)

    def _add_contract(self, S, docx_path):
        s = S()
        c = Contract(user_id=1, file_name="原合同.docx", stored_path=str(docx_path),
                     parsed_text=PARSED, status="completed")
        s.add(c)
        s.commit()
        cid = c.id
        s.close()
        return cid

    def _user(self):
        u = mock.MagicMock()
        u.id = 1
        u.role = "uploader"
        return u

    def _fake_rag(self):
        m = types.ModuleType("ai.rag")
        m.search_knowledge = mock.MagicMock(return_value=None)
        m.search_similar_templates = mock.MagicMock(return_value=None)
        return m

    def _revise(self, S, cid, **kw):
        s = S()
        try:
            with mock.patch.dict(sys.modules, {"ai.rag": self._fake_rag()}):
                return contracts.revise_contract_clause(
                    cid, contracts.ReviseRequest(**kw), db=s, current_user=self._user())
        finally:
            s.close()

    def _download(self, S, cid):
        s = S()
        try:
            with mock.patch.object(contracts, "UPLOAD_DIR", str(self.uploads)):
                return contracts.download_revised_docx(
                    cid, BackgroundTasks(), db=s, current_user=self._user())
        finally:
            s.close()

    # ── 1. 比对来源：无 AuditRecord 锚点 → _locate_clause 兜底 ──
    def test_comparison_source_gets_anchor_via_locate_clause(self):
        tmp, eng, S = self._setup()
        try:
            with tempfile.TemporaryDirectory() as fdir:
                docx = Path(fdir) / "原合同.docx"
                self._make_docx(docx)
                cid = self._add_contract(S, docx)
                result = {"revised_clause": "第二条 验收标准与验收方式：本合同项下的验收采用人工审核并出具验收书的方式。",
                          "constraints": [], "legal_basis": [], "remaining_risks": [], "explanation": "ok"}
                with mock.patch.object(contracts, "revise_clause", return_value=dict(result)):
                    res = self._revise(S, cid, clause_text=MATCHED, instruction="补全验收方式",
                                       scope="clause", clause_key=CMP_KEY, clause_no="")
                self.assertEqual(res["code"], 0)

                s = S()
                rev = s.query(ClauseRevision).order_by(ClauseRevision.id.desc()).first()
                anchor, no = rev.original_clause_text, rev.clause_no
                s.close()
                self.assertTrue(anchor, "比对来源应通过 _locate_clause 生成原文锚点")
                self.assertIn(MATCHED[:10], anchor)
                self.assertEqual(no, "2", "兜底时应同步 clause_no")

                # 且 DOCX 能正确替换，其它条款不被破坏
                resp = self._download(S, cid)
                texts = [p.text for p in Document(resp.path).paragraphs if p.text.strip()]
                os.remove(resp.path)
                self.assertEqual(len(texts), 3)
                self.assertIn("人工审核并出具验收书", texts[1])
                self.assertEqual(texts[0], DOC_PARAS[0])
                self.assertEqual(texts[2], DOC_PARAS[2])
        finally:
            eng.dispose()
            tmp.cleanup()

    # ── 2. 比对会话多轮：沿用上一轮锚点，DOCX 用最新结果 ──
    def test_comparison_multi_round_reuses_previous_anchor(self):
        tmp, eng, S = self._setup()
        try:
            with tempfile.TemporaryDirectory() as fdir:
                docx = Path(fdir) / "原合同.docx"
                self._make_docx(docx)
                cid = self._add_contract(S, docx)
                r1 = {"revised_clause": "第二条 验收标准：采用人工审核并出具验收书。",
                      "constraints": [], "legal_basis": [], "remaining_risks": [], "explanation": "r1"}
                with mock.patch.object(contracts, "revise_clause", return_value=dict(r1)):
                    self._revise(S, cid, clause_text=MATCHED, instruction="补全", scope="clause",
                                 clause_key=CMP_KEY, clause_no="")
                # 第二轮：clause_text 已是上一轮修订稿（原文中已不存在）
                r2 = {"revised_clause": "第二条 验收标准：采用人工审核并出具验收书，且由第三方复核。",
                      "constraints": [], "legal_basis": [], "remaining_risks": [], "explanation": "r2"}
                with mock.patch.object(contracts, "revise_clause", return_value=dict(r2)):
                    self._revise(S, cid, clause_text=r1["revised_clause"], instruction="再加第三方复核",
                                 scope="clause", clause_key=CMP_KEY, clause_no="2")

                s = S()
                rev2 = s.query(ClauseRevision).order_by(ClauseRevision.id.desc()).first()
                anchor2 = rev2.original_clause_text
                s.close()
                self.assertTrue(anchor2, "第二轮应沿用上一轮锚点")

                resp = self._download(S, cid)
                texts = [p.text for p in Document(resp.path).paragraphs if p.text.strip()]
                os.remove(resp.path)
                self.assertIn("第三方复核", texts[1], "DOCX 应采用最新一轮结果")
                self.assertEqual(texts[2], DOC_PARAS[2])
        finally:
            eng.dispose()
            tmp.cleanup()

    # ── 3. 无锚点修订 → 明确 400，不静默交付漏改 DOCX ──
    def test_download_refuses_when_revision_has_no_anchor(self):
        tmp, eng, S = self._setup()
        try:
            with tempfile.TemporaryDirectory() as fdir:
                docx = Path(fdir) / "原合同.docx"
                self._make_docx(docx)
                cid = self._add_contract(S, docx)
                s = S()
                # 一条可定位 + 一条无锚点（模拟历史数据/定位失败）
                s.add(ClauseRevision(contract_id=cid, scope="clause", clause_key="1",
                                     clause_text=RISK_CLAUSE, instruction="改",
                                     revised_clause="第三条 违约责任与违约金上限为合同总价的10%。",
                                     original_clause_text=RISK_ANCHOR))
                s.add(ClauseRevision(contract_id=cid, scope="clause", clause_key="__cmp__X",
                                     clause_text="定位不到的条款", instruction="改",
                                     revised_clause="改后内容", original_clause_text=None))
                s.commit()
                s.close()
                with self.assertRaises(HTTPException) as ctx:
                    self._download(S, cid)
                self.assertEqual(ctx.exception.status_code, 400)
                self.assertIn("无法在合同原文中定位", str(ctx.exception.detail))
        finally:
            eng.dispose()
            tmp.cleanup()

    # ── 4. 风险来源锚点仍取自 AuditRecord（兜底不改原逻辑） ──
    def test_risk_source_anchor_still_from_audit_record(self):
        tmp, eng, S = self._setup()
        try:
            with tempfile.TemporaryDirectory() as fdir:
                docx = Path(fdir) / "原合同.docx"
                self._make_docx(docx)
                cid = self._add_contract(S, docx)
                s = S()
                rec = AuditRecord(contract_id=cid, audit_batch="b1", risk_type="R02",
                                  risk_level="high", clause_text=RISK_CLAUSE,
                                  detection_method="evidence", result_status="valid",
                                  clause_position={"original_text": RISK_ANCHOR, "clause_no": 3})
                s.add(rec)
                s.commit()
                rid = rec.id
                s.close()
                result = {"revised_clause": "第三条 违约责任与违约金上限为合同总价的20%。",
                          "constraints": [], "legal_basis": [], "remaining_risks": [], "explanation": "ok"}
                with mock.patch.object(contracts, "revise_clause", return_value=dict(result)):
                    self._revise(S, cid, clause_text=RISK_CLAUSE, instruction="下调上限",
                                 scope="clause", clause_key=str(rid), clause_no="3")
                s = S()
                rev = s.query(ClauseRevision).order_by(ClauseRevision.id.desc()).first()
                anchor = rev.original_clause_text
                s.close()
                self.assertEqual(anchor, RISK_ANCHOR, "风险来源必须仍用审核阶段锚点")
        finally:
            eng.dispose()
            tmp.cleanup()


    # ── 5. add_clause 专项会话：scope=clause 且不带原文锚点（防被替换路径重复应用） ──
    def test_add_clause_session_is_scope_clause_without_anchor(self):
        tmp, eng, S = self._setup()
        try:
            with tempfile.TemporaryDirectory() as fdir:
                docx = Path(fdir) / "原合同.docx"
                self._make_docx(docx)
                cid = self._add_contract(S, docx)
                gen = {"clause_text": "第四条 付款方式：合同签订后10日内支付30%。",
                       "explanation": "新增", "legal_basis": []}
                with mock.patch.object(contracts, "generate_clause", return_value=dict(gen)):
                    res = self._revise(S, cid, clause_text="", instruction="新增付款条款",
                                       scope="clause", clause_key="__cmp__付款条款",
                                       operation="add_clause",
                                       position={"append": True, "hint": "追加到合同末尾"})
                self.assertEqual(res["code"], 0)

                s = S()
                rev = s.query(ClauseRevision).order_by(ClauseRevision.id.desc()).first()
                scope, op, key, anchor = rev.scope, rev.operation, rev.clause_key, rev.original_clause_text
                s.close()
                self.assertEqual((scope, op, key), ("clause", "add_clause", "__cmp__付款条款"))
                self.assertIsNone(anchor, "新增条款不得带原文锚点，否则会被替换路径重复应用")

                # 且仍能正常插入 DOCX 并编号顺延
                resp = self._download(S, cid)
                texts = [p.text for p in Document(resp.path).paragraphs if p.text.strip()]
                os.remove(resp.path)
                self.assertEqual(len(texts), 4)
                self.assertTrue(texts[3].startswith("四、"), texts[3])
        finally:
            eng.dispose()
            tmp.cleanup()

    # ── 6. 总体会话：专项会话汇总必须真的进入发给 LLM 的 prompt ──
    def test_overview_instruction_reaches_llm_prompt(self):
        tmp, eng, S = self._setup()
        try:
            with tempfile.TemporaryDirectory() as fdir:
                docx = Path(fdir) / "原合同.docx"
                self._make_docx(docx)
                cid = self._add_contract(S, docx)
                digest = ("【当前修改依据】以下为各专项会话的修改记录，请结合这些要求从整个合同的角度统一修改：\n"
                          "【风险专项修改记录】\n- 风险 R08（中风险）：\n  用户最终要求：补充验收程序\n"
                          "【条款比对专项修改记录】\n- 条款比对·验收标准：\n  用户最终要求：补全第三方复核\n\n"
                          "【用户修改要求】\n结合前面讨论过的问题，统一修改这份合同。")
                captured = []

                def fake_chat(prompt, temperature=0.1):
                    captured.append(prompt)
                    return json.dumps({"constraints": [], "legal_basis": [], "revised_clause": "整体修订后的正文",
                                       "explanation": "ok", "verified": True, "remaining_risks": []},
                                      ensure_ascii=False)

                s = S()
                try:
                    with mock.patch.dict(sys.modules, {"ai.rag": self._fake_rag()}), \
                         mock.patch.object(reviser.llm_client, "chat", side_effect=fake_chat):
                        contracts.revise_contract_clause(
                            cid, contracts.ReviseRequest(clause_text=PARSED, instruction=digest,
                                                         scope="overview", clause_key="__overview__"),
                            db=s, current_user=self._user())
                finally:
                    s.close()

                # Leader + Follower 驱动实际修订，必须含专项会话汇总
                # （Self-QA 按 reviser 既有设计只接收原条款/修订后条款/Leader 约束，不含用户指令）
                self.assertGreaterEqual(len(captured), 2)
                for prompt in captured[:2]:
                    self.assertIn("【风险专项修改记录】", prompt)
                    self.assertIn("【条款比对专项修改记录】", prompt)
                    self.assertIn("统一修改这份合同", prompt)
        finally:
            eng.dispose()
            tmp.cleanup()


if __name__ == "__main__":
    unittest.main()