"""阶段 1：PDF 合同 → 修订版 DOCX 测试。

覆盖用户要求的 A–H 全部场景：

  A. 重建无损        parsed_text → 中间 DOCX → 逐字回读一致
  B. 标题识别        条款结构（第X条 / X、）不被破坏
  C. 空白折叠匹配    LLM 把换行写成空格的 clause_text 能定位
  D. 多命中拒绝      同文出现两次 → 不建立锚点
  E. 短文本拒绝      低于可靠阈值 → 不建立锚点
  F. 真实 PDF 95     4/4 风险条款建立可靠锚点（R03/R04/R05/R06）
  G. 真实修订 E2E    PDF 95 → 修改 → adopted → GET /revised-docx → 读回验证
  H. R09 E2E         PDF 中选择「第七条后」新增条款 → 正确插入 + 后续编号顺延

以及边界回归：
  - PDF 锚点无法可靠定位时导出返回明确 400，**不产生任何文件**
  - DOCX 合同导出行为不变（原有链路零回归）
  - 导出是只读的：不新建/不修改任何 ClauseRevision，不改 AuditRecord 风险结果
  - 中间 DOCX 生成在临时目录，不污染 backend/data

运行（backend 目录下）：
    python -m pytest tests/test_pdf_intermediate_docx.py -v
"""
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

_BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))
# PaddleOCR / 大模型都不需要：本文件只碰解析文本、锚点算法与 DOCX 读写
os.environ.setdefault("SECRET_KEY", "test-secret-key-not-weak-123456")
os.environ.setdefault("DEEPSEEK_API_KEY", "sk-test-not-real")

from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from fastapi import HTTPException  # noqa: E402

from database import Base  # noqa: E402
from models.contract import Contract  # noqa: E402
from models.audit_record import AuditRecord  # noqa: E402
from models.clause_revision import ClauseRevision  # noqa: E402
from services.intermediate_docx import build_from_parsed_text, collapse_ws  # noqa: E402
from services.pdf_anchor import (  # noqa: E402
    build_anchor, MIN_ANCHOR_LEN, norm_with_map,
)
from services.pdf_anchor_backfill import backfill_pdf_anchors, is_pdf_like  # noqa: E402
from services.docx_reviser import build_revised_docx, _norm, _heading_num  # noqa: E402
import api.contracts as contracts  # noqa: E402

REAL_DB = _BACKEND_DIR / "contract.db"

# ── 合成 fixture：刻意复刻真实 PDF 的结构（目录区 + 正文区、跨行条款）──
# 目录区**必须与真实合同同形**：真实 PDF 的目录里「八、其他」下面完整列出了
# 1.考察 2.驻场 3.奖励 4.保密 5.联络 6.知识产权 7.廉洁自律 8.成果文件 —— 正是这段
# 目录条目把目录的「八、其他…4. 保密」与正文的「4. 保密…委托人申明…」拉开了几十行，
# 远超 MAX_GAP=2；锚点算法据此才会选中正文而不是目录（已用真实 PDF 95 交叉验证）。
PARSED = "\n".join([
    "海南省建设工程造价咨询合同（示范文本 HNFJ—2026—0001）",
    "合同编号: HN-ZJ-2026-0812",
    "八、其他",                     # ← 目录区（与正文区同名，用于多命中测试）
    "1. 考察及相关费用",
    "2. 驻场",
    "3. 奖励",
    "4. 保密",
    "5. 联络",
    "6. 知识产权",
    "7. 廉洁自律",
    "8. 成果文件",
    "第三部分 专用条款",
    "八、其他",
    "4. 保密",                      # ← 正文区
    "委托人申明的保密事项和期限：项目所有造价信息及商业资料，保密期限永久。",
    "咨询人申明的保密事项和期限：本合同内容及服务过程信息，保密期限永久有效。",
    "第三人申明的保密事项和期限：无。",
    "5. 联络",
    "6. 知识产权",
    "6.2 咨询人为履行本合同约定而编制的成果文件，其著作权属于双方共有。",
    "6.3 双方将履行本合同形成的有关成果文件用于企业宣传、申报奖项以及接受上级主",
    "管部门的检查须遵守以下约定：需经对方书面同意后方可使用。",
    "六、合同变更、解除与终止",
    "1. 合同变更",
    "2. 合同解除",
    "2.2 双方约定解除合同的条件还包括：委托人有权随时单方解除本合同，只需提前 15",
    "日书面通知咨询人，按已完成工作量结算咨询费，不承担其他违约责任。",
    "七、争议解决",
    "4. 仲裁或诉讼",
    "（2）向咨询人所在地海口市龙华区人民法院提起诉讼。",
    "九、补充条款：",
])

# LLM 证据文本（换行 → 空格；标题/编号/正文被合并）—— 与真实数据同形
CLAUSE_R03 = "六、合同变更、解除与终止 2. 合同解除 2.2 双方约定解除合同的条件还包括：委托人有权随时单方解除本合同，只需提前 15 日书面通知咨询人，按已完成工作量结算咨询费，不承担其他违约责任。"
CLAUSE_R04 = "向咨询人所在地海口市龙华区人民法院提起诉讼。"
CLAUSE_R05 = "八、其他 4. 保密 委托人申明的保密事项和期限：项目所有造价信息及商业资料，保密期限永久。咨询人申明的保密事项和期限：本合同内容及服务过程信息，保密期限永久有效。第三人申明的保密事项和期限：无。"
CLAUSE_R06 = "八、其他 6. 知识产权 6.2 咨询人为履行本合同约定而编制的成果文件，其著作权属于双方共有。6.3 双方将履行本合同形成的有关成果文件用于企业宣传、申报奖项以及接受上级主管部门的检查须遵守以下约定：需经对方书面同意后方可使用。"


class TestIntermediateDocxRebuild(unittest.TestCase):
    """A + B：重建无损、标题结构不被破坏。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.out = os.path.join(self.tmp.name, "intermediate.docx")

    def test_A_rebuild_is_verbatim_lossless(self):
        """A. 逐行回读必须与 parsed_text 逐字一致（含空行）。"""
        n = build_from_parsed_text(PARSED, self.out)
        doc = Document(self.out)
        self.assertEqual(n, len(PARSED.split("\n")))
        self.assertEqual("\n".join(p.text for p in doc.paragraphs), PARSED)

    def test_A_trailing_newline_is_not_silently_trimmed(self):
        """尾随换行必须变成空段落，不允许偷偷 trim。"""
        text = "第一条 甲\n第二条 乙\n"
        build_from_parsed_text(text, self.out)
        paras = [p.text for p in Document(self.out).paragraphs]
        self.assertEqual(paras, ["第一条 甲", "第二条 乙", ""])
        self.assertEqual("\n".join(paras), text)

    def test_A_leading_and_interior_blank_lines_preserved(self):
        text = "\n第一条 甲\n\n第二条 乙"
        build_from_parsed_text(text, self.out)
        paras = [p.text for p in Document(self.out).paragraphs]
        self.assertEqual(paras, ["", "第一条 甲", "", "第二条 乙"])
        self.assertEqual("\n".join(paras), text)

    def test_A_empty_parsed_text_yields_single_empty_paragraph(self):
        build_from_parsed_text("", self.out)
        paras = [p.text for p in Document(self.out).paragraphs]
        self.assertEqual(paras, [""])
        self.assertEqual("\n".join(paras), "")

    def test_B_clause_headings_survive_rebuild(self):
        """B. 「第X条 / X、」结构在中间 DOCX 中仍可被识别（docx_reviser 依赖它）。"""
        build_from_parsed_text(PARSED, self.out)
        doc = Document(self.out)
        nums = sorted({n for n in (_heading_num(p.text) for p in doc.paragraphs) if n is not None})
        # 六、七、八、九 + 4、5、6、1、2 等次级编号
        for expected in (6, 7, 8, 9):
            self.assertIn(expected, nums, f"顶层编号 {expected} 未被识别")
        # 与 _parse_headings 的口径一致
        from api.contracts import _parse_headings
        src_nums = sorted({h["num"] for h in _parse_headings(PARSED)})
        self.assertEqual(src_nums, nums)

    def test_normalized_fulltext_matches_parsed_text(self):
        """折叠口径下中间 DOCX 全文 == parsed_text（docx_reviser 的定位前提）。"""
        build_from_parsed_text(PARSED, self.out)
        doc_all = "\n".join(p.text for p in Document(self.out).paragraphs)
        self.assertEqual(_norm(doc_all), _norm(PARSED))


class TestAnchorAlgorithm(unittest.TestCase):
    """C + D + E：空白折叠匹配、多命中拒绝、短文本拒绝。"""

    def test_C_whitespace_folded_cross_line_clause_is_located(self):
        """C. LLM 把换行写成空格 + 合并标题/编号/正文，仍能定位。"""
        built = build_anchor(PARSED, CLAUSE_R03)
        self.assertIsNotNone(built, "跨行合并的 clause_text 必须能定位")
        text, start, end = built
        self.assertEqual(PARSED[start:end], text, "锚点必须是 parsed_text 的真实连续子串")
        self.assertIn("双方约定解除合同的条件还包括", text)
        self.assertGreaterEqual(len(text), MIN_ANCHOR_LEN)

    def test_C_anchor_is_real_substring_not_llm_rewrite(self):
        """锚点原文必须来自 parsed_text，不能是 LLM 改写后的文本。"""
        built = build_anchor(PARSED, CLAUSE_R05)
        self.assertIsNotNone(built)
        text, _, _ = built
        self.assertIn(text, PARSED)
        # LLM 文本里被拼接的目录行不应出现在锚点里
        self.assertNotIn("八、其他 4. 保密 委托人", text)
        # 锚点覆盖正文三句
        self.assertIn("委托人申明的保密事项和期限", text)
        self.assertIn("第三人申明的保密事项和期限", text)

    def test_C_multi_sentence_merges_only_when_adjacent(self):
        """句级命中相邻（只隔换行）才合并；跨区域的分散命中必须拒绝。"""
        built = build_anchor(PARSED, CLAUSE_R05)
        self.assertIsNotNone(built)
        _, start, end = built
        span = PARSED[start:end]
        # 三句都在同一连续区间里
        for frag in ("委托人申明的保密事项和期限", "咨询人申明的保密事项和期限", "第三人申明的保密事项和期限"):
            self.assertIn(frag, span)

    def test_C_prefix_only_hit_still_located(self):
        """LLM 尾部被改写（如 6.3 那句）时，仍能用唯一命中的前缀定位。"""
        built = build_anchor(PARSED, CLAUSE_R06)
        self.assertIsNotNone(built, "R06 这类尾部被改写的条款必须能定位")
        text, _, _ = built
        self.assertIn("6.2 咨询人为履行本合同约定而编制的成果文件", text)
        self.assertIn(text, PARSED)

    def test_D_multiple_hits_are_rejected(self):
        """D. 同一文本在 parsed_text 中出现两次 → 拒绝建立锚点（不选第一个）。"""
        parsed = "合同解除。\n中间其它内容很多很多很多。\n合同解除。"
        self.assertIsNone(build_anchor(parsed, "合同解除。"))
        # 即便 LLM 文本更长，只要其唯一可命中的片段是重复的，也必须拒绝
        self.assertIsNone(build_anchor(parsed, "合同解除。 补充说明"))

    def test_D_duplicate_in_real_fixture_rejected(self):
        """fixture 里「4. 保密」「八、其他」都出现两次；只给编号的条款必须拒绝。"""
        self.assertIsNone(build_anchor(PARSED, "4. 保密"))
        self.assertIsNone(build_anchor(PARSED, "八、其他"))

    def test_E_short_text_is_rejected(self):
        """E. 低于可靠阈值不建立锚点（即便文本在正文中唯一存在）。"""
        short = "甲" * (MIN_ANCHOR_LEN - 1)      # 7 字
        parsed_uniq = "前缀内容" + short + "后缀内容"
        self.assertIsNone(build_anchor(parsed_uniq, short),
                          f"{len(short)} 字 < MIN_ANCHOR_LEN={MIN_ANCHOR_LEN} 必须拒绝")
        self.assertIsNone(build_anchor(PARSED, "短"))
        # 恰好达到阈值且唯一 → 允许
        parsed = "前缀内容甲乙丙丁戊己庚辛后缀"
        built = build_anchor(parsed, "甲乙丙丁戊己庚辛")
        self.assertIsNotNone(built, "恰好 8 字且唯一应被接受")
        self.assertEqual(built[0], "甲乙丙丁戊己庚辛")
        # 7 字（差 1 字）在同一段正文里 → 拒绝
        self.assertIsNone(build_anchor(parsed, "甲乙丙丁戊己庚"))

    def test_E_partial_prefix_truncation_stops_at_min_len(self):
        """前缀逐级缩短到阈值以下必须停止（不得退化成 4 字近似命中）。"""
        # "唯一但很短"的片段不存在 → 拒绝，而不是命中更短的子串
        parsed = "完全无关的内容甲甲甲甲"
        self.assertIsNone(build_anchor(parsed, "甲乙丙丁戊己庚辛壬癸"))

    def test_anchor_never_returns_whitespace_only(self):
        self.assertIsNone(build_anchor(PARSED, "       "))
        self.assertIsNone(build_anchor("", CLAUSE_R04))
        self.assertIsNone(build_anchor(PARSED, ""))

    def test_norm_map_is_position_accurate(self):
        """归一化下标映射必须能精确还原原文区间（含全角空格与多空行）。"""
        text = "甲  乙\n\n丙\u3000丁"
        norm, cmap = norm_with_map(text)
        self.assertEqual(len(norm), len(cmap))
        for i, ch in enumerate(norm):
            if ch == " ":
                self.assertTrue(text[cmap[i]].isspace())
            else:
                self.assertEqual(text[cmap[i]], ch)


@unittest.skipUnless(REAL_DB.is_file(), "真实库 contract.db 不存在")
class TestRealPdf95Anchors(unittest.TestCase):
    """F：真实 PDF 合同 95 的 4 条风险必须全部建立可靠锚点。"""

    @classmethod
    def setUpClass(cls):
        import sqlite3
        con = sqlite3.connect(f"file:{REAL_DB}?mode=ro", uri=True)
        con.row_factory = sqlite3.Row
        try:
            row = con.execute("select parsed_text, stored_path from contracts where id=95").fetchone()
            if row is None:
                raise unittest.SkipTest("真实库里没有合同 95")
            cls.parsed = row["parsed_text"] or ""
            cls.stored_path = row["stored_path"]
            cls.records = [
                dict(r) for r in con.execute(
                    "select id, risk_type, clause_text from audit_records "
                    "where contract_id=95 and result_status='valid' order by id")
            ]
        finally:
            con.close()

    def test_F_synthetic_fixture_baseline_4_of_4(self):
        """先在合成 fixture 上锁定 4/4（不依赖真实库内容漂移）。"""
        got = {}
        for rt, ct in (("R03", CLAUSE_R03), ("R04", CLAUSE_R04),
                       ("R05", CLAUSE_R05), ("R06", CLAUSE_R06)):
            built = build_anchor(PARSED, ct)
            got[rt] = built is not None
        self.assertEqual(got, {"R03": True, "R04": True, "R05": True, "R06": True})

    def test_F_real_pdf95_all_risks_get_reliable_anchor(self):
        """F. 真实 PDF 95：目标 4/4（本轮前为 1/4）。"""
        self.assertTrue(self.stored_path.lower().endswith(".pdf"), "合同 95 应为 PDF")
        self.assertEqual(len(self.records), 4, "合同 95 应有 4 条有效风险")
        result = {}
        for r in self.records:
            built = build_anchor(self.parsed, r["clause_text"] or "")
            result[r["risk_type"]] = built
        for rt in ("R03", "R04", "R05", "R06"):
            self.assertIsNotNone(result.get(rt), f"{rt} 未能建立可靠锚点")
            text, start, end = result[rt]
            self.assertEqual(self.parsed[start:end], text)
            self.assertGreaterEqual(len(text), MIN_ANCHOR_LEN)
        self.assertEqual(sum(1 for v in result.values() if v), 4, "应为 4/4")

    def test_F_real_pdf95_anchor_is_verbatim_original(self):
        """锚点必须是 parsed_text 的真实连续子串，且不含 LLM 拼接出来的错配文本。"""
        for r in self.records:
            built = build_anchor(self.parsed, r["clause_text"] or "")
            self.assertIsNotNone(built)
            text, start, end = built
            self.assertIn(text, self.parsed)
            self.assertEqual(self.parsed.find(text), start)


class TestAnchorBackfill(unittest.TestCase):
    """后置锚点回填：只写 clause_position，不动风险判定，幂等。"""

    def setUp(self):
        # Windows 上临时目录里的 sqlite 文件必须先归还连接池才能删除
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        eng = create_engine(f"sqlite:///{Path(self.tmp.name) / 't.db'}", connect_args={"timeout": 30})
        self.engine = eng
        Base.metadata.create_all(eng)
        self.S = sessionmaker(bind=eng)

    def tearDown(self):
        self.engine.dispose()

    def _seed(self, stored_name="c.pdf"):
        s = self.S()
        c = Contract(user_id=1, file_name="测试合同", stored_path=f"/tmp/{stored_name}",
                     parsed_text=PARSED, contract_type="建设工程合同", status="completed")
        s.add(c)
        s.commit()
        cid = c.id
        rows = [
            ("R03", CLAUSE_R03, "high", "违约金比例过高"),
            ("R04", CLAUSE_R04, "medium", "管辖不利"),
            ("R05", CLAUSE_R05, "high", "保密期限永久"),
            ("R06", CLAUSE_R06, "low", "知识产权不清"),
        ]
        for rt, ct, lv, reason in rows:
            s.add(AuditRecord(contract_id=cid, audit_batch="b1", risk_type=rt, risk_level=lv,
                              clause_text=ct, reason=reason, suggestion="建议修改",
                              detection_method="rule", confidence=0.8, result_status="valid",
                              clause_position=None))
        s.commit()
        s.close()
        return cid

    def _fresh(self, cid):
        s = self.S()
        c = s.query(Contract).filter(Contract.id == cid).first()
        recs = s.query(AuditRecord).filter(AuditRecord.contract_id == cid).order_by(AuditRecord.id).all()
        return s, c, recs

    def test_backfill_fills_all_four_anchors(self):
        cid = self._seed()
        s, c, _ = self._fresh(cid)
        stats = backfill_pdf_anchors(s, c, audit_batch="b1")
        s.close()
        self.assertEqual(stats["total"], 4)
        self.assertEqual(stats["anchored"], 4)
        self.assertEqual(stats["unreliable"], 0)

        s, _, recs = self._fresh(cid)
        for r in recs:
            pos = r.clause_position or {}
            self.assertTrue((pos.get("original_text") or "").strip(), f"{r.risk_type} 锚点为空")
            self.assertIn(pos["original_text"], PARSED)
            self.assertEqual(PARSED[pos["start"]:pos["end"]], pos["original_text"])
        s.close()

    def test_backfill_does_not_touch_risk_verdict_fields(self):
        """只写 clause_position；风险判定相关字段必须逐字不变。"""
        cid = self._seed()
        s, c, before = self._fresh(cid)
        snapshot = [(r.risk_type, r.risk_level, r.reason, r.suggestion, r.confidence,
                     r.detection_method, r.result_status) for r in before]
        backfill_pdf_anchors(s, c, audit_batch="b1")
        s.close()

        s, _, after = self._fresh(cid)
        now = [(r.risk_type, r.risk_level, r.reason, r.suggestion, r.confidence,
                r.detection_method, r.result_status) for r in after]
        self.assertEqual(snapshot, now)
        s.close()

    def test_backfill_is_idempotent(self):
        cid = self._seed()
        s, c, _ = self._fresh(cid)
        first = backfill_pdf_anchors(s, c, audit_batch="b1")
        s.close()
        s, c, recs1 = self._fresh(cid)
        anchors1 = [json.dumps(r.clause_position, sort_keys=True, ensure_ascii=False) for r in recs1]
        s.close()

        s, c, _ = self._fresh(cid)
        second = backfill_pdf_anchors(s, c, audit_batch="b1")
        s.close()
        s, _, recs2 = self._fresh(cid)
        anchors2 = [json.dumps(r.clause_position, sort_keys=True, ensure_ascii=False) for r in recs2]
        s.close()

        self.assertEqual(anchors1, anchors2, "重复回填不得改变锚点")
        self.assertEqual(first["anchored"], 4)
        self.assertEqual(second["anchored"], 0)
        self.assertEqual(second["unchanged"], 4)

    def test_backfill_leaves_unreliable_records_without_anchor(self):
        cid = self._seed()
        s = self.S()
        s.add(AuditRecord(contract_id=cid, audit_batch="b1", risk_type="R09", risk_level="medium",
                          clause_text="完全无法在正文中定位的杜撰条款内容一段。",
                          detection_method="rule", result_status="valid", clause_position=None))
        s.commit()
        s.close()

        s, c, _ = self._fresh(cid)
        stats = backfill_pdf_anchors(s, c, audit_batch="b1")
        s.close()
        self.assertEqual(stats["unreliable"], 1)
        self.assertEqual(stats["anchored"], 4)

        s, _, recs = self._fresh(cid)
        r09 = next(r for r in recs if r.risk_type == "R09")
        self.assertIsNone(r09.clause_position, "不可靠的记录不应被写入锚点")
        s.close()

    def test_is_pdf_like(self):
        class C:
            def __init__(self, p):
                self.stored_path = p
        self.assertTrue(is_pdf_like(C("/data/x.pdf")))
        self.assertTrue(is_pdf_like(C("/data/x.PNG")))
        self.assertFalse(is_pdf_like(C("/data/x.docx")))
        self.assertFalse(is_pdf_like(C("")))
        self.assertFalse(is_pdf_like(C(None)))


def _mid_docx(parsed_text, tmpdir):
    path = os.path.join(tmpdir, "intermediate.docx")
    build_from_parsed_text(parsed_text, path)
    return path


class TestBuildRevisedDocxOnIntermediate(unittest.TestCase):
    """G + H（单元层）：中间 DOCX 上直接跑现有 build_revised_docx。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.src = _mid_docx(PARSED, self.tmp.name)

    @staticmethod
    def _rev(rid, clause_text, revised_clause, anchor, operation="replace", position=None,
             scope="clause", adopted=True):
        class R:
            pass
        r = R()
        r.id = rid
        r.scope = scope
        r.operation = operation
        r.clause_text = clause_text
        r.revised_clause = revised_clause
        r.original_clause_text = anchor
        r.position = position
        r.adopted = adopted
        return r

    def test_G_replace_all_four_and_verify_in_place(self):
        """G（单元）：4 条替换必须原地生效，且目标原文消失、其它内容保留。"""
        plans = [
            ("R03", CLAUSE_R03, "【修订R03】重新约定了单方解除条件与补偿安排。"),
            ("R04", CLAUSE_R04, "【修订R04】向委托人所在地人民法院提起诉讼。"),
            ("R05", CLAUSE_R05, "【修订R05】保密期限为合同终止后三年。"),
            ("R06", CLAUSE_R06, "【修订R06】成果文件著作权归委托人所有。"),
        ]
        revs, anchors = [], {}
        for i, (rt, ct, new) in enumerate(plans, start=1):
            built = build_anchor(PARSED, ct)
            self.assertIsNotNone(built, f"{rt} 应能建立锚点")
            anchors[rt] = built[0]
            revs.append(self._rev(i, ct, new, built[0]))

        out = os.path.join(self.tmp.name, "revised.docx")
        applied, skipped = build_revised_docx(self.src, revs, out)
        self.assertEqual((applied, skipped), (4, 0))

        doc = Document(out)
        text = _norm("\n".join(p.text for p in doc.paragraphs))
        for rt, _, new in plans:
            self.assertIn(new, text, f"{rt} 的修订文本未出现")
        # 原文其余部分保留
        self.assertIn("海南省建设工程造价咨询合同", text)
        self.assertIn("第三部分 专用条款", text)
        self.assertIn("本合同内容及服务过程信息".replace("本合同内容及服务过程信息", "九、补充条款："), text)

    def test_H_add_clause_after_seventh_inserts_and_renumbers(self):
        """H（单元）：在「第七条之后」新增 → 插入正确 + 后续顶层编号顺延。"""
        rev = self._rev(99, "", "【新增】不可抗力：因不可抗力不能履行的，部分或全部免除责任。",
                        None, operation="add_clause", position={"anchor": "七", "hint": "第七条之后"})
        out = os.path.join(self.tmp.name, "added.docx")
        applied, skipped = build_revised_docx(self.src, [rev], out)
        self.assertEqual(skipped, 0)
        self.assertEqual(applied, 1)

        doc = Document(out)
        paras = [p.text for p in doc.paragraphs]
        text = "\n".join(paras)
        self.assertIn("【新增】不可抗力", text)

        i_new = next(i for i, t in enumerate(paras) if "【新增】不可抗力" in t)
        i_seven = next(i for i, t in enumerate(paras) if t.strip().startswith("七、争议解决"))
        # 语义：插在「第七条」这一条款**全部内容之后**、下一个顶层条款之前。
        # 因此必然排在第七条标题之后，并恰好占住顺延后空出的那个位置。
        self.assertGreater(i_new, i_seven, "新增条款应在第七条标题之后")
        i_next_top = next(
            (i for i, t in enumerate(paras)
             if i > i_seven and re.match(r"^[一二三四五六七八九十]+、", t.strip())),
            None,
        )
        self.assertEqual(i_new, i_next_top, "新增条款应恰好占住顺延后空出的位置")

        # 后续顶层编号顺延：原「八、其他」「九、补充条款」→「九、其他」「十、补充条款」
        tops = [t.strip() for t in paras if re.match(r"^[一二三四五六七八九十]+、", t.strip())]
        self.assertIn("九、其他", tops)
        self.assertTrue(any(t.startswith("十、补充条款") for t in tops), f"应顺延出「十、补充条款」，实际 {tops}")
        self.assertNotIn("八、其他", tops, "原第八条应已被顺延为第九条")
        self.assertFalse(any(t.startswith("九、补充条款") for t in tops), "「九、补充条款」应已顺延为「十、」")

    def test_H_add_clause_append_at_end(self):
        rev = self._rev(98, "", "【新增末尾】补充条款：本补充条款自双方签字之日起生效。",
                        None, operation="add_clause", position={"append": True})
        out = os.path.join(self.tmp.name, "appended.docx")
        applied, _ = build_revised_docx(self.src, [rev], out)
        self.assertEqual(applied, 1)
        paras = [p.text for p in Document(out).paragraphs]
        self.assertIn("【新增末尾】补充条款", paras[-1])

    def test_H_duplicate_clause_numbers_insert_after_first_match_documented(self):
        """**已知既有行为**（本轮未改 `docx_reviser`，故不修）：

        `_insert_paragraph_level` 用 `_heading_num(p.text) == anchor_num` 找**第一个**匹配段。
        真实合同里同一顶层编号会在目录/正文（甚至通用条款/专用条款两部分）重复出现 ——
        实测 22 个真实 DOCX 中 12 个存在重复顶层编号 —— 因此新增条款会落在
        **第一个**编号匹配的条款之后，而不是用户心里那个同名章节之后。

        本用例把这个确定性行为钉住，避免后续误以为是 PDF 链路新引入的问题；
        要真正解决需按"用户在位置选择器里选中的那一条"定位，属独立任务。
        """
        rev = self._rev(97, "", "【新增-重复编号】条款正文。", None,
                        operation="add_clause", position={"anchor": "八", "hint": "第八条之后"})
        out = os.path.join(self.tmp.name, "dupnum.docx")
        applied, _ = build_revised_docx(self.src, [rev], out)
        self.assertEqual(applied, 1)
        paras = [p.text for p in Document(out).paragraphs]
        i_new = next(i for i, t in enumerate(paras) if "【新增-重复编号】" in t)
        # 确定性事实（与 docx_reviser 现有实现一致，本轮未改动该模块）：
        #   ① 新条款编号 = anchor + 1；
        #   ② 它排在第一个编号匹配的顶层条款**之后**；
        #   ③ 文档只增加 1 段（插入而非覆盖）。
        # fixture 里第一个匹配「八、」的是目录区那一条（第 2 段），故新条款落在其后；
        # 正文区那条「八、其他」由现有实现决定是否顺延 —— 这里只钉住上述不变式。
        eights = [i for i, t in enumerate(paras) if re.match(r"^八、", t.strip())]
        self.assertGreaterEqual(len(eights), 1, "fixture 应至少有一处「八、」")
        self.assertGreater(i_new, eights[0])
        self.assertEqual(_heading_num(paras[i_new]), 9, "新条款编号应为 anchor+1")
        self.assertEqual(len(paras), len(PARSED.split("\n")) + 1, "只应新增 1 段")


class TestPdfExportEndpoint(unittest.TestCase):
    """G：PDF 合同的 /revised-docx 端点（真实函数调用 + 真实文件读写）。"""

    def setUp(self):
        # Windows 上临时目录里的 sqlite 文件必须先归还连接池才能删除
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        eng = create_engine(f"sqlite:///{Path(self.tmp.name) / 't.db'}", connect_args={"timeout": 30})
        self.engine = eng
        Base.metadata.create_all(eng)
        self.S = sessionmaker(bind=eng)
        self.uploads = Path(self.tmp.name) / "uploads"
        self.uploads.mkdir(exist_ok=True)

    def tearDown(self):
        self.engine.dispose()

    def _user(self, uid=1, role="uploader"):
        class U:
            pass
        u = U()
        u.id = uid
        u.role = role
        return u

    def _seed_pdf_contract(self, with_records=True, parsed_text=PARSED):
        pdf = Path(self.tmp.name) / "src.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake original\n")
        s = self.S()
        c = Contract(user_id=1, file_name="测试合同", stored_path=str(pdf),
                     parsed_text=parsed_text, contract_type="建设工程合同", status="completed")
        s.add(c)
        s.commit()
        cid = c.id
        if with_records:
            for rt, ct, lv in (("R03", CLAUSE_R03, "high"), ("R04", CLAUSE_R04, "medium"),
                               ("R05", CLAUSE_R05, "high"), ("R06", CLAUSE_R06, "low")):
                s.add(AuditRecord(contract_id=cid, audit_batch="b1", risk_type=rt, risk_level=lv,
                                  clause_text=ct, detection_method="rule", result_status="valid",
                                  clause_position=None))
            s.commit()
        s.close()
        return cid, pdf

    def _add_revision(self, cid, clause_text, revised_clause, operation="replace",
                      anchor=None, position=None):
        s = self.S()
        r = ClauseRevision(contract_id=cid, scope="clause", operation=operation,
                           clause_key="1", clause_text=clause_text,
                           original_clause_text=anchor, revised_clause=revised_clause,
                           instruction="测试", adopted=True, position=position)
        s.add(r)
        s.commit()
        rid = r.id
        s.close()
        return rid

    def _download(self, cid):
        s = self.S()
        try:
            with mock.patch.object(contracts, "UPLOAD_DIR", str(self.uploads)):
                return contracts.download_revised_docx(
                    cid, background_tasks=mock.MagicMock(), db=s, current_user=self._user())
        finally:
            s.close()

    def test_G_pdf_contract_exports_revised_docx(self):
        """G. PDF 合同（锚点为空、需实时重算）能成功导出修订版 DOCX。"""
        cid, pdf = self._seed_pdf_contract()
        for ct, new in ((CLAUSE_R03, "【修订R03】重新约定了单方解除条件与补偿安排。"),
                        (CLAUSE_R04, "【修订R04】向委托人所在地人民法院提起诉讼。"),
                        (CLAUSE_R05, "【修订R05】保密期限为合同终止后三年。"),
                        (CLAUSE_R06, "【修订R06】成果文件著作权归委托人所有。")):
            self._add_revision(cid, ct, new)

        resp = self._download(cid)
        doc = Document(resp.path)
        text = _norm("\n".join(p.text for p in doc.paragraphs))
        for marker in ("【修订R03】", "【修订R04】", "【修订R05】", "【修订R06】"):
            self.assertIn(marker, text, f"{marker} 未写入修订版")
        self.assertIn("海南省建设工程造价咨询合同", text, "原文其它内容必须保留")
        self.assertEqual(doc.paragraphs[0].text, PARSED.split("\n")[0])
        os.remove(resp.path)

    def test_G_original_pdf_is_not_modified(self):
        cid, pdf = self._seed_pdf_contract()
        before = pdf.read_bytes()
        self._add_revision(cid, CLAUSE_R04, "【修订R04】改后文本。")
        resp = self._download(cid)
        os.remove(resp.path)
        self.assertEqual(pdf.read_bytes(), before, "原始 PDF 绝不能被改写")

    def test_G_export_does_not_create_or_modify_revisions(self):
        """导出是只读的：ClauseRevision 数量与内容、AuditRecord 风险结果都不变。"""
        cid, _ = self._seed_pdf_contract()
        self._add_revision(cid, CLAUSE_R04, "【修订R04】改后文本。")
        s = self.S()
        before_revs = [(r.id, r.clause_text, r.revised_clause, r.original_clause_text, r.adopted)
                       for r in s.query(ClauseRevision).filter(ClauseRevision.contract_id == cid).all()]
        before_recs = [(r.id, r.risk_type, r.risk_level, r.clause_text, r.clause_position)
                       for r in s.query(AuditRecord).filter(AuditRecord.contract_id == cid).all()]
        s.close()

        resp = self._download(cid)
        os.remove(resp.path)

        s = self.S()
        after_revs = [(r.id, r.clause_text, r.revised_clause, r.original_clause_text, r.adopted)
                      for r in s.query(ClauseRevision).filter(ClauseRevision.contract_id == cid).all()]
        after_recs = [(r.id, r.risk_type, r.risk_level, r.clause_text, r.clause_position)
                      for r in s.query(AuditRecord).filter(AuditRecord.contract_id == cid).all()]
        s.close()
        self.assertEqual(before_revs, after_revs, "导出不得新建/修改 ClauseRevision")
        self.assertEqual(before_recs, after_recs, "导出不得修改 AuditRecord（含风险结果）")

    def test_unreliable_anchor_returns_400_and_writes_no_file(self):
        """锚点无法可靠定位 → 明确 400，且不产生任何文件（不返回"看起来成功"的 DOCX）。"""
        cid, _ = self._seed_pdf_contract()
        self._add_revision(cid, "完全杜撰、绝不可能出现在合同正文里的条款内容一段。", "【无效修订】")
        with self.assertRaises(HTTPException) as ctx:
            self._download(cid)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("无法在原文中可靠定位", ctx.exception.detail)
        # 不产生残留修订版文件
        self.assertEqual([p.name for p in self.uploads.glob("revised_*.docx")], [])

    def test_no_revisions_returns_400(self):
        cid, _ = self._seed_pdf_contract()
        with self.assertRaises(HTTPException) as ctx:
            self._download(cid)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("还没有任何条款修改", ctx.exception.detail)

    def test_pdf_without_parsed_text_returns_400(self):
        cid, _ = self._seed_pdf_contract(with_records=False, parsed_text="")
        self._add_revision(cid, CLAUSE_R04, "【修订】改后。")
        with self.assertRaises(HTTPException) as ctx:
            self._download(cid)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("没有可用的解析正文", ctx.exception.detail)

    def test_missing_source_file_and_no_parsed_text_returns_400(self):
        """原件已被清理且没有解析正文 → 明确 400，而不是生成空修订版或抛 500。"""
        cid, pdf = self._seed_pdf_contract(with_records=False, parsed_text="")
        self._add_revision(cid, CLAUSE_R04, "【修订】改后。")
        pdf.unlink()
        with self.assertRaises(HTTPException) as ctx:
            self._download(cid)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("没有可用的解析正文", ctx.exception.detail)
        self.assertEqual([p.name for p in self.uploads.glob("revised_*.docx")], [])

    def test_missing_source_file_with_parsed_text_falls_back_to_intermediate(self):
        """原件被清理但解析文本仍在 → 用中间 DOCX 兜底继续导出（不弱于改动前行为）。"""
        cid, pdf = self._seed_pdf_contract()
        self._add_revision(cid, CLAUSE_R04, "【修订R04】兜底导出。")
        pdf.unlink()
        resp = self._download(cid)
        text = _norm("\n".join(p.text for p in Document(resp.path).paragraphs))
        self.assertIn("【修订R04】兜底导出。", text)
        os.remove(resp.path)

    def test_intermediate_docx_is_not_written_into_upload_dir(self):
        """中间 DOCX 必须生成在临时目录，不污染 backend/data（UPLOAD_DIR）。"""
        cid, _ = self._seed_pdf_contract()
        self._add_revision(cid, CLAUSE_R04, "【修订R04】改后文本。")
        resp = self._download(cid)
        os.remove(resp.path)
        leftovers = [p.name for p in self.uploads.iterdir()]
        self.assertEqual(leftovers, [], f"UPLOAD_DIR 不应有残留：{leftovers}")

    def test_docx_contract_still_uses_original_file(self):
        """回归：DOCX 合同仍直接使用原文件，行为不变。"""
        docx_path = Path(self.tmp.name) / "src.docx"
        d = Document()
        for line in PARSED.split("\n"):
            d.add_paragraph(line)
        d.save(str(docx_path))

        s = self.S()
        c = Contract(user_id=1, file_name="测试合同", stored_path=str(docx_path),
                     parsed_text=PARSED, contract_type="建设工程合同", status="completed")
        s.add(c)
        s.commit()
        cid = c.id
        s.close()

        built = build_anchor(PARSED, CLAUSE_R04)
        self._add_revision(cid, CLAUSE_R04, "【DOCX修订】改后文本。", anchor=built[0])
        resp = self._download(cid)
        text = _norm("\n".join(p.text for p in Document(resp.path).paragraphs))
        self.assertIn("【DOCX修订】", text)
        os.remove(resp.path)


@unittest.skipUnless(REAL_DB.is_file(), "真实库 contract.db 不存在")
class TestRealPdf95EndToEnd(unittest.TestCase):
    """G（真实内容）：用真实 PDF 合同 95 的 parsed_text + 真实风险 clause_text 跑完整导出。

    刻意**不写真实库**：只从真实库**只读**取内容，再灌进临时库跑端点，
    这样既能验证真实数据的可修订性，又不会污染任何历史数据。
    """

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        eng = create_engine(f"sqlite:///{Path(self.tmp.name) / 't.db'}", connect_args={"timeout": 30})
        self.engine = eng
        Base.metadata.create_all(eng)
        self.S = sessionmaker(bind=eng)
        self.uploads = Path(self.tmp.name) / "uploads"
        self.uploads.mkdir(exist_ok=True)

    def tearDown(self):
        self.engine.dispose()

    def _user(self):
        class U:
            pass
        u = U()
        u.id = 1
        u.role = "uploader"
        return u

    def _download(self, cid):
        s = self.S()
        try:
            with mock.patch.object(contracts, "UPLOAD_DIR", str(self.uploads)):
                return contracts.download_revised_docx(
                    cid, background_tasks=mock.MagicMock(), db=s, current_user=self._user())
        finally:
            s.close()

    def test_G_real_pdf95_full_chain(self):
        import sqlite3
        con = sqlite3.connect(f"file:{REAL_DB}?mode=ro", uri=True)
        con.row_factory = sqlite3.Row
        try:
            src = con.execute("select parsed_text from contracts where id=95").fetchone()
            if src is None:
                self.skipTest("真实库里没有合同 95")
            parsed = src["parsed_text"] or ""
            recs = [dict(r) for r in con.execute(
                "select risk_type, clause_text from audit_records "
                "where contract_id=95 and result_status='valid' order by id")]
        finally:
            con.close()
        self.assertEqual(len(recs), 4, "合同 95 应有 4 条有效风险")

        # 真实 PDF 原件（只读复制到临时目录，确认导出不会改它）
        pdf = Path(self.tmp.name) / "real95.pdf"
        real_pdf = REAL_DB.parent / "data"
        matches = sorted(real_pdf.glob("*.pdf"))
        if matches:
            import shutil
            shutil.copyfile(matches[0], pdf)
        else:
            pdf.write_bytes(b"%PDF-1.4 placeholder\n")
        before = pdf.read_bytes()

        s = self.S()
        c = Contract(user_id=1, file_name="测试合同文件", stored_path=str(pdf),
                     parsed_text=parsed, contract_type="建设工程合同", status="completed")
        s.add(c)
        s.commit()
        cid = c.id
        s.close()

        # 真实风险 → 逐条模拟「用户已确认采用此版」的修订（锚点留空，由端点实时重算）
        expected = {}
        for i, r in enumerate(recs, start=1):
            new_text = f"【真实E2E修订{r['risk_type']}】该条款已按审核建议重新拟定。"
            expected[r["risk_type"]] = new_text
            s = self.S()
            s.add(ClauseRevision(contract_id=cid, scope="clause", operation="replace",
                                 clause_key=str(i), clause_text=r["clause_text"] or "",
                                 original_clause_text=None, revised_clause=new_text,
                                 instruction="真实 E2E", adopted=True, position=None))
            s.commit()
            s.close()

        resp = self._download(cid)
        doc = Document(resp.path)
        text = _norm("\n".join(p.text for p in doc.paragraphs))
        try:
            for rt, new_text in expected.items():
                self.assertIn(new_text, text, f"{rt} 的修订未写入真实 PDF 的修订版")
            # 原文主体保留（首行 + 末尾附录）
            self.assertIn(_norm(parsed.split("\n")[0]), text)
            self.assertIn("附录 E 咨询人服务团队成员表", text)
        finally:
            os.remove(resp.path)

        self.assertEqual(pdf.read_bytes(), before, "真实 PDF 原件不得被改写")

        # 导出未新建任何修订记录
        s = self.S()
        self.assertEqual(
            s.query(ClauseRevision).filter(ClauseRevision.contract_id == cid).count(), len(recs))
        s.close()


class TestFrontendExportGate(unittest.TestCase):
    """前端可导出判据：PDF 放开，图片仍不支持（源码级契约）。"""

    SRC = _BACKEND_DIR.parent / "frontend" / "src" / "composables" / "workspaceLogic.js"

    def _src(self):
        self.assertTrue(self.SRC.is_file(), f"找不到 {self.SRC}")
        return self.SRC.read_text(encoding="utf-8")

    def test_pdf_is_exportable_docx_and_image_are_handled(self):
        src = self._src()
        self.assertIn("EXPORTABLE_SOURCE_EXTS", src)
        self.assertRegex(src, r"EXPORTABLE_SOURCE_EXTS\s*=\s*\[\s*'docx'\s*,\s*'pdf'\s*\]")
        # 旧文案必须消失
        self.assertNotIn("仅 DOCX 原始合同支持导出修订版", src)
        self.assertNotIn("当前上传的不是 Word 文档", src)
        # 新文案说明统一输出 DOCX
        self.assertIn("暂不支持导出修订版", src)

    def test_pdf_case_reaches_ready_state(self):
        """用 node 跑真实函数，确认 pdf 不再是 not_docx、图片仍是 not_docx。"""
        import subprocess
        script = (
            "import('./workspaceLogic.js').then(m => {"
            "  const f = m.docxStateFor;"
            "  const out = {"
            "    pdf: f({storedPath:'/d/a.pdf', exportableCount:2, blockerCount:0}).key,"
            "    docx: f({storedPath:'/d/a.docx', exportableCount:2, blockerCount:0}).key,"
            "    png: f({storedPath:'/d/a.png', exportableCount:2, blockerCount:0}).key,"
            "    jpg: f({storedPath:'/d/a.jpg', exportableCount:2, blockerCount:0}).key,"
            "    noext: f({storedPath:'/d/a', exportableCount:2, blockerCount:0}).key,"
            "    pdfNoRev: f({storedPath:'/d/a.pdf', exportableCount:0, blockerCount:0}).key,"
            "    docxBlocked: f({storedPath:'/d/a.docx', exportableCount:2, blockerCount:1}).key,"
            "    nameIgnored: f({storedPath:'/d/a.pdf', fileName:'x.docx', exportableCount:1, blockerCount:0}).key"
            "  };"
            "  console.log(JSON.stringify(out));"
            "})"
        )
        proc = subprocess.run(
            ["node", "--input-type=module", "-e", script],
            cwd=str(self.SRC.parent), capture_output=True, text=True, timeout=60,
        )
        self.assertEqual(proc.returncode, 0, f"node 执行失败: {proc.stderr[:400]}")
        got = json.loads(proc.stdout.strip().splitlines()[-1])
        self.assertEqual(got["pdf"], "ready", "PDF 合同必须可导出")
        self.assertEqual(got["docx"], "ready")
        self.assertEqual(got["png"], "not_docx", "图片本轮仍不支持")
        self.assertEqual(got["jpg"], "not_docx", "图片本轮仍不支持")
        self.assertEqual(got["noext"], "not_docx")
        self.assertEqual(got["pdfNoRev"], "none")
        self.assertEqual(got["docxBlocked"], "blocked")
        self.assertEqual(got["nameIgnored"], "ready", "必须以 stored_path 判断格式")


if __name__ == "__main__":
    unittest.main()
