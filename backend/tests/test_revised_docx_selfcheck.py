"""S-1 回归：DOCX / 普通 PDF / 扫描 PDF / JPG / PNG 上「替换已有条款 + R09 新增条款」
同一批导出修订版 DOCX 必须 200，且导出自检不再出现假阴性。

## 问题（根因，见本轮汇报）

`docx_reviser._insert_paragraph_level` / `_insert_inline` 在插入新增条款时会**顺延后续标题编号**
（编号 > anchor_num 的标题 +N）。若被替换条款位于插入点之后，**修订文本自身的行首编号也会被
系统合法改写**：

    R09 在「第三条」插入 → 已替换的「第五条 保密…」被顺延成「第六条 保密…」

而 `api.contracts._verify_revised_docx` 用「修订文本必须逐字出现」做自检（只对非 DOCX 路径生效），
于是把"已写入"误判成"未写入" → 导出被拦成 400。实测扫描 PDF / JPG / PNG 全部命中。

## 修复（A 方案，最小改动）

自检在**抹掉标题编号**（`_HEADING_RE` 命中的同一批 token，即顺延逻辑唯一会改动的东西）
之后再做一次比对；且**仅当本次导出确实包含会被插入的 add_clause 修订时**才启用该口径
（纯替换导出仍逐字比对，行为与修复前完全一致）。正文仍必须逐字出现，因此
「正文没写进去」照旧会被拦下 —— 安全网没有被关闭，也没有被放松成任意文本都通过。

运行（backend 目录下）：
    python -m pytest tests/test_revised_docx_selfcheck.py -v
"""
import difflib
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

_BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))
os.environ.setdefault("SECRET_KEY", "test-secret-key-not-weak-123456")
os.environ.setdefault("DEEPSEEK_API_KEY", "sk-test-not-real")

from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from fastapi import HTTPException  # noqa: E402

from database import Base  # noqa: E402
from models.contract import Contract  # noqa: E402
from models.clause_revision import ClauseRevision  # noqa: E402
from services.docx_reviser import (  # noqa: E402
    _norm, _heading_num, _normalize_heading_numbers, build_revised_docx,
)
from api.contracts import _verify_revised_docx  # noqa: E402
import api.contracts as contracts  # noqa: E402

# ── fixture：刻意复刻 S-1 的真实结构 ───────────────────────────────────────
# 替换发生在「第五条」（编号 5），R09 插在「第三条」（anchor_num=3）之后
# → 顺延逻辑必然把已替换条款的行首「第五条」改成「第六条」。
S1_PARSED = "\n".join([
    "建设工程施工合同",                                  # 0
    "第一条 工程概况",                                   # 1
    "工程名称：海口江东新区某项目。",                     # 2
    "第二条 合同价款",                                   # 3
    "合同价款为人民币壹佰万元整。",                       # 4
    "第三条 违约责任",                                   # 5 ← R09 插入锚点（paragraph_index=5）
    "任何一方违约应向对方支付违约金。",                   # 6
    "第四条 争议解决",                                   # 7 ← 顺延后应变为「第五条 争议解决」
    "争议提交有管辖权的人民法院解决。",                   # 8
    "第五条 保密",                                       # 9 ← 被替换条款的标题
    "双方对合同内容承担永久保密义务。",                   # 10 ← 被替换条款正文
    "第六条 附则",                                       # 11
    "本合同自双方签字之日起生效。",                       # 12
])

# 替换：LLM 证据文本（换行→空格，标题与正文拼成一行）与修订稿
S1_CLAUSE_TEXT = "第五条 保密 双方对合同内容承担永久保密义务。"
S1_REVISED = ("第五条 保密 一、保密信息的范围 本条款所称保密信息是指双方在履行本合同过程中"
              "知悉的对方商业秘密与技术资料，保密期限为合同终止后三年。")
S1_REVISED_BODY = "保密信息的范围"          # 只出现在修订稿里，用于断言"确实写进去"
S1_ADD_CLAUSE = ("不可抗力条款：一、不可抗力的情形 本合同所称不可抗力，是指不能预见、不能避免"
                 "且不能克服的客观情况；二、通知义务 受影响一方应在三日内书面通知对方。")
S1_ADD_BODY = "不能预见、不能避免"
S1_POSITION = {"anchor": "三", "hint": "第三条（违约责任）之后", "target_text": "第三条 违约责任",
               "paragraph_index": 5}
S1_TARGET_INDEX = 5
S1_EXPECTED_INSERT_INDEX = 7      # 第四条 争议解决（原第 7 段）之前


def _kind_marker(kind: str) -> str:
    return f"[{kind}]"


class _Rev:
    """轻量修订对象（docx_reviser / _verify_revised_docx 只用 getattr 取字段）。"""

    def __init__(self, rid, revised_clause, *, operation="replace", anchor=None,
                 clause_text="", position=None, scope="clause"):
        self.id = rid
        self.scope = scope
        self.operation = operation
        self.clause_text = clause_text
        self.revised_clause = revised_clause
        self.original_clause_text = anchor
        self.position = position
        self.adopted = True


def _write_docx(path, parsed: str):
    doc = Document()
    for line in parsed.split("\n"):
        doc.add_paragraph(line)
    doc.save(path)
    return path


def _paras(path):
    return [p.text for p in Document(str(path)).paragraphs]


def _unexplained_changes(baseline, final, *, inserted_body, replace_body):
    """返回"无法解释"的变更段落（除 新增段 / 替换段 / 多段兜底清空 / 编号顺延 之外）。"""
    ins = _normalize_heading_numbers(_norm(inserted_body))
    rep = _normalize_heading_numbers(_norm(replace_body))
    sm = difflib.SequenceMatcher(a=baseline, b=final, autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        base_block = [_norm(t) for t in baseline[i1:i2] if _norm(t)]
        for j in range(j1, j2):
            t = _norm(final[j])
            if not t:
                continue                                   # 多段兜底被清空的段落
            tn = _normalize_heading_numbers(t)
            if ins and ins in tn:
                continue                                   # R09 新增段
            if rep and rep in tn:
                continue                                   # 替换后的段落
            if any(tn == _normalize_heading_numbers(b) for b in base_block):
                continue                                   # 仅标题编号顺延（正文逐字相同）
            out.append((j, t))
    return out


class TestHeadingNumberNormalization(unittest.TestCase):
    """归一化只抹编号，不抹正文。"""

    def test_only_heading_numbers_are_normalized(self):
        pairs = [
            ("第五条 保密 一、范围", "第六条 保密 一、范围"),
            ("第五条 保密 一、范围", "第五条 保密 二、范围"),
            ("第 5 条 付款", "第 9 条 付款"),
            ("八、其他 4. 保密 正文", "九、其他 4. 保密 正文"),
        ]
        for a, b in pairs:
            self.assertEqual(_normalize_heading_numbers(a), _normalize_heading_numbers(b),
                             f"编号差异应被归一化：{a!r} vs {b!r}")

    def test_body_difference_is_not_normalized(self):
        self.assertNotEqual(_normalize_heading_numbers("第五条 保密 期限三年"),
                            _normalize_heading_numbers("第六条 保密 期限五年"))
        self.assertNotEqual(_normalize_heading_numbers("第五条 保密"),
                            _normalize_heading_numbers("第六条 竞业限制"))

    def test_normalized_text_keeps_body_verbatim(self):
        text = "第五条 保密 一、保密信息的范围 本条款所称保密信息是指"
        got = _normalize_heading_numbers(text)
        for frag in ("保密", "保密信息的范围", "本条款所称保密信息是指"):
            self.assertIn(frag, got)
        # 编号被换成等长的占位符：长度不变，变化只发生在编号本身
        self.assertEqual(len(got), len(text))
        self.assertNotIn("第五条", got)
        self.assertNotIn("一、", got)

    def test_empty_and_none_are_safe(self):
        self.assertEqual(_normalize_heading_numbers(""), "")
        self.assertEqual(_normalize_heading_numbers(None), None)

    def test_insertion_logic_still_sees_headings(self):
        """顺延逻辑依赖的 `_heading_num` 不受本函数影响。"""
        self.assertEqual(_heading_num("第六条 保密"), 6)
        self.assertEqual(_heading_num("四、争议解决"), 4)


class TestSelfCheckStrictness(unittest.TestCase):
    """自检的"从严"方向必须保住：真漏改仍然要被抓到。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self.tmp.cleanup)
        self.out = os.path.join(self.tmp.name, "out.docx")
        # 结果文件：替换后的条款已被系统顺延成「第六条」，R09 新增条款也已写入
        self.add_text = "四、不可抗力条款：一、不可抗力的情形 不能预见、不能避免且不能克服。"
        self.anchor = "第六条 保密"          # 非空锚点，替换修订才会进入自检的 expected
        _write_docx(self.out, "\n".join([
            "第三条 违约责任",
            self.add_text,
            "第六条 保密 一、保密信息的范围 保密期限三年。",
        ]))

    def test_verbatim_text_passes(self):
        revs = [_Rev(1, "第六条 保密 一、保密信息的范围 保密期限三年。", anchor=self.anchor)]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [])

    def test_s1_renumbered_leading_number_passes_when_add_clause_present(self):
        """S-1：修订文本正文逐字在文件里，只有行首编号被顺延 → 必须通过。"""
        revs = [
            # 修订稿原文是「第五条 …」，导出时被 R09 顺延成「第六条 …」
            _Rev(1, "第五条 保密 一、保密信息的范围 保密期限三年。", anchor=self.anchor),
            _Rev(2, self.add_text, operation="add_clause", position=S1_POSITION),
        ]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [],
                         "编号由系统自身顺延造成的差异不得判为「未写入」（S-1）")

    def test_replace_only_export_keeps_verbatim_check(self):
        """纯替换导出（无 add_clause）→ 仍逐字比对，行为与修复前一致。"""
        revs = [_Rev(1, "第五条 保密 一、保密信息的范围 保密期限三年。", anchor=self.anchor)]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [1],
                         "没有新增条款就不存在编号顺延，编号不同必须判为未写入")

    def test_body_missing_is_still_caught_even_with_add_clause(self):
        """正文没写进去（哪怕编号对得上）→ 仍然失败：安全网没被放松。"""
        revs = [
            _Rev(1, "第六条 保密 一、竞业限制的范围 完全不同的正文内容。", anchor=self.anchor),
            _Rev(2, self.add_text, operation="add_clause", position=S1_POSITION),
        ]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [1])

    def test_completely_absent_text_is_still_caught(self):
        revs = [
            _Rev(1, "第九条 完全不存在的条款正文，一个字都没写进结果文件。", anchor=self.anchor),
            _Rev(2, self.add_text, operation="add_clause", position=S1_POSITION),
        ]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [1])

    def test_too_short_text_cannot_pass_on_number_alone(self):
        """归一化后过短的文本不算命中（避免"只对上一个编号"就通过）。"""
        _write_docx(self.out, "第六条 保密\n" + self.add_text)
        revs = [
            _Rev(1, "第五条 保密", anchor=self.anchor),      # 归一化后仅 2 字
            _Rev(2, self.add_text, operation="add_clause", position=S1_POSITION),
        ]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [1])

    def test_anchorless_revision_is_not_reported_here(self):
        """无锚点修订由端点上层 anchorless 拦截，自检本身不重复报（既有语义）。"""
        revs = [_Rev(1, "【任意】", anchor="")]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [])

    def test_whitespace_only_revised_clause_is_ignored(self):
        revs = [_Rev(1, "   ")]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [])

    def test_add_clause_text_that_is_missing_is_caught(self):
        revs = [_Rev(2, "四、这段新增条款根本没进文件。", operation="add_clause", position=S1_POSITION)]
        self.assertEqual(_verify_revised_docx(self.out, revs, {}), [2])


class TestReplacePlusAddClauseExportAllFormats(unittest.TestCase):
    """S-1 端点级回归：五种输入格式「替换 + R09 新增」同批导出必须 200 且内容正确。"""

    SOURCE_KINDS = ("docx", "pdf", "scanned_pdf", "jpg", "png")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        eng = create_engine(f"sqlite:///{Path(self.tmp.name) / 't.db'}", connect_args={"timeout": 30})
        self.engine = eng
        Base.metadata.create_all(eng)
        self.S = sessionmaker(bind=eng)
        self.uploads = Path(self.tmp.name) / "uploads"
        self.uploads.mkdir(exist_ok=True)

    def tearDown(self):
        self.engine.dispose()

    # ── 基础设施 ──────────────────────────────────────────────────────────
    def _user(self):
        class U:
            pass
        u = U()
        u.id = 1
        u.role = "uploader"
        return u

    def _source(self, kind: str) -> Path:
        """五种输入的原件：DOCX 用真实 DOCX（走原件路径），其余按扩展名建占位原件。

        非 DOCX 导出一律由 `parsed_text` 重建中间 DOCX，**不会重新读原件**
        （与既有 `test_pdf_intermediate_docx` 的 fixture 口径一致）。
        """
        if kind == "docx":
            return Path(_write_docx(str(Path(self.tmp.name) / "src.docx"), S1_PARSED))
        ext = {"pdf": ".pdf", "scanned_pdf": ".pdf", "jpg": ".jpg", "png": ".png"}[kind]
        p = Path(self.tmp.name) / f"src_{kind}{ext}"
        if ext == ".pdf":
            p.write_bytes(b"%PDF-1.4 fake source (content is rebuilt from parsed_text)\n")
        else:
            from PIL import Image
            Image.new("RGB", (32, 32), (255, 255, 255)).save(str(p))
        return p

    def _seed(self, kind: str):
        """建合同 + 一条替换修订 + 一条 add_clause 修订（与真实链路同形）。"""
        src = self._source(kind)
        s = self.S()
        c = Contract(user_id=1, file_name=f"测试合同{_kind_marker(kind)}", stored_path=str(src),
                     parsed_text=S1_PARSED, contract_type="建设工程合同", status="completed")
        s.add(c)
        s.commit()
        cid = c.id
        # 替换：DOCX 路径用库内锚点（原件路径不重算），其余格式锚点留空由端点实时重算。
        # 锚点都取「标题 + 正文」的整条条款（与真实审核落库的 clause_position 同形），
        # 因此两条都会走 `_find_paragraph_indices` 的多段兜底，DOCX 与非 DOCX 结构一致。
        anchor = S1_CLAUSE_TEXT if kind == "docx" else None
        s.add(ClauseRevision(contract_id=cid, scope="clause", operation="replace",
                             clause_key="1", clause_text=S1_CLAUSE_TEXT,
                             original_clause_text=anchor, revised_clause=S1_REVISED,
                             instruction="S-1 回归", adopted=True, position=None))
        s.add(ClauseRevision(contract_id=cid, scope="clause", operation="add_clause",
                             clause_key="__r09__", clause_text="",
                             original_clause_text=None, revised_clause=S1_ADD_CLAUSE,
                             instruction="R09 新增不可抗力", adopted=True, position=S1_POSITION))
        s.commit()
        revs = s.query(ClauseRevision).filter(ClauseRevision.contract_id == cid).all()
        s.close()
        return cid, src, revs

    def _download(self, cid):
        s = self.S()
        try:
            with mock.patch.object(contracts, "UPLOAD_DIR", str(self.uploads)):
                return contracts.download_revised_docx(
                    cid, background_tasks=mock.MagicMock(), db=s, current_user=self._user())
        finally:
            s.close()

    def _assert_export_ok(self, kind: str):
        cid, src, _ = self._seed(kind)
        # ① 核心：S-1 修复前这里会抛 400「修订版生成自检未通过」
        resp = self._download(cid)
        try:
            final = _paras(resp.path)
            text = _norm("\n".join(final))

            # ② 替换内容确实写入（编号被顺延，正文逐字在）
            self.assertIn(S1_REVISED_BODY, text, f"{kind}: 替换后的条款正文未写入")
            # ③ R09 新增内容确实写入
            self.assertIn(S1_ADD_BODY, text, f"{kind}: R09 新增条款未写入")

            # ④ R09 插入位置正确（在第三个标题之后、下一个更高编号标题之前）
            i_new = next(i for i, t in enumerate(final) if S1_ADD_BODY in t)
            self.assertEqual(i_new, S1_EXPECTED_INSERT_INDEX, f"{kind}: 插入位置错误")
            self.assertGreater(i_new, S1_TARGET_INDEX, f"{kind}: 必须插在所选目标条款之后")
            self.assertEqual(_heading_num(final[i_new]), 4, f"{kind}: 新条款编号应为 4")

            # ⑤ 编号顺延正确：1,2,3,4(新增),5,6,7 单调递增，且原文标题按规则整体 +1
            nums = [_heading_num(t) for t in final if _heading_num(t) is not None]
            self.assertEqual(nums, sorted(nums), f"{kind}: 编号必须单调递增：{nums}")
            self.assertEqual(nums[:7], [1, 2, 3, 4, 5, 6, 7], f"{kind}: 顶层编号应为 1..7：{nums}")
            self.assertTrue(any(t.strip().startswith("第五条 争议解决") for t in final),
                            f"{kind}: 原「第四条 争议解决」应顺延为「第五条」")
            self.assertTrue(any("第六条 保密" in t for t in final),
                            f"{kind}: 被替换条款的行首编号应随系统顺延为「第六条」")
            self.assertTrue(any(t.strip().startswith("第七条 附则") for t in final),
                            f"{kind}: 原「第六条 附则」应顺延为「第七条」")

            # ⑥ 除 替换 / 新增 / 编号顺延 / 多段兜底清空 外，没有 unexplained changes
            unexplained = _unexplained_changes(
                S1_PARSED.split("\n"), final,
                inserted_body=S1_ADD_CLAUSE, replace_body=S1_REVISED)
            self.assertEqual(unexplained, [], f"{kind}: 出现未预期改动：{unexplained}")

            # ⑦ 原文其它内容保留
            self.assertIn("合同价款为人民币壹佰万元整。", text, f"{kind}: 原文其它内容被破坏")
            self.assertIn("本合同自双方签字之日起生效。", text, f"{kind}: 原文其它内容被破坏")

            # ⑧ 自检本身在产物上返回空（不再假阴性）。
            #   非 DOCX 必须带上端点实时重算的锚点，否则替换修订会被当成"无锚点"跳过，
            #   自检就退化成只查新增条款 —— 这里刻意用重算锚点把**替换**也纳入核对。
            s = self.S()
            revs = s.query(ClauseRevision).filter(ClauseRevision.contract_id == cid).all()
            s.close()
            if kind == "docx":
                overrides = {}
            else:
                overrides, unreliable = contracts._resolve_pdf_anchors(revs, S1_PARSED)
                self.assertEqual(unreliable, [], f"{kind}: 替换修订未能重算出可靠锚点")
                self.assertEqual(len(overrides), 1, f"{kind}: 应恰好重算出 1 条替换锚点")
            self.assertEqual(_verify_revised_docx(resp.path, revs, overrides), [],
                             f"{kind}: 自检仍判定有修订未写入")
            # ⑨ DOCX 可重新打开 = 可编辑
            self.assertGreaterEqual(len(final), len(S1_PARSED.split("\n")))
        finally:
            if os.path.exists(getattr(resp, "path", "")):
                os.remove(resp.path)

    # ── 五种格式 ──────────────────────────────────────────────────────────
    def test_docx_replace_plus_r09_exports_ok(self):
        self._assert_export_ok("docx")

    def test_text_pdf_replace_plus_r09_exports_ok(self):
        self._assert_export_ok("pdf")

    def test_scanned_pdf_replace_plus_r09_exports_ok(self):
        self._assert_export_ok("scanned_pdf")

    def test_jpg_replace_plus_r09_exports_ok(self):
        self._assert_export_ok("jpg")

    def test_png_replace_plus_r09_exports_ok(self):
        self._assert_export_ok("png")

    # ── 端点级"仍然从严"的正向对照 ────────────────────────────────────────
    def test_unresolvable_add_clause_position_fails_loudly(self):
        """生成异常（新增位置无法可靠定位）→ 明确 400，且不留下任何输出文件。"""
        cid, _, _ = self._seed("png")
        s = self.S()
        add = (s.query(ClauseRevision)
               .filter(ClauseRevision.contract_id == cid,
                       ClauseRevision.operation == "add_clause").first())
        add.position = {"anchor": "三", "target_text": "完全不存在的标题文本"}
        s.commit()
        s.close()
        with self.assertRaises(HTTPException) as ctx:
            self._download(cid)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertTrue(("未找到" in ctx.exception.detail) or ("无法唯一" in ctx.exception.detail),
                        ctx.exception.detail)
        self.assertEqual([p.name for p in self.uploads.glob("revised_*.docx")], [],
                         "失败时不得留下「看起来成功」的文件")

    def test_anchorless_replace_still_400(self):
        """无锚点 → 仍然明确 400（安全网①：不允许静默漏改）。"""
        cid, _, _ = self._seed("png")
        s = self.S()
        rep = (s.query(ClauseRevision)
               .filter(ClauseRevision.contract_id == cid,
                       ClauseRevision.operation == "replace").first())
        rep.clause_text = "完全杜撰、绝不可能出现在合同正文里的内容一段。"   # 让 pdf_anchor 也算不出锚点
        s.commit()
        s.close()
        with self.assertRaises(HTTPException) as ctx:
            self._download(cid)
        self.assertEqual(ctx.exception.status_code, 400)
        self.assertIn("无法在原文中可靠定位", ctx.exception.detail)
        self.assertEqual([p.name for p in self.uploads.glob("revised_*.docx")], [])


if __name__ == "__main__":
    unittest.main()
